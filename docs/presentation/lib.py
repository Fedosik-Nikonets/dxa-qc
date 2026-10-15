# Хелперы для сборки презентации на шаблоне ЛЦТ 2026 (python-pptx).
# Шрифт не задаём: текст наследует шрифт темы шаблона (Montserrat).
import copy
import re
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageFont

PINK = "FF0053"      # accent1
PINK_L = "FFD6E3"    # accent2
PINK2 = "FC3777"     # accent3
LAV = "8A83D1"       # accent4
PURP = "520977"      # accent5
DPURP = "2D1451"     # accent6
INK = "1C1D22"       # dk2
GREY = "E5E7E9"      # lt2
WHITE = "FFFFFF"
MUTED = "5F5A6E"
LAV_L = "F1ECF7"     # светлая плашка внутри белых карточек
LINE_L = "D9D2E6"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

def _load_measure_font():
    for path in ("Montserrat-Bold.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
                 "/Library/Fonts/Montserrat-Bold.ttf",
                 "/System/Library/Fonts/Supplemental/Verdana Bold.ttf"):
        try:
            return ImageFont.truetype(path, 1000)
        except OSError:
            continue
    return None


_measure_font = _load_measure_font()


def text_width_in(text, pt):
    """Ширина строки (дюймы) жирным шрифтом — по ней считается ширина плашки заголовка.
    DejaVu Sans Bold и Verdana Bold близки по ширине к Montserrat Bold."""
    if _measure_font is None:
        return len(text) * 0.72 * pt / 72
    return _measure_font.getlength(text) / 1000 * pt / 72


def rgb(h):
    return RGBColor.from_string(h)


def rrect(s, x, y, w, h, fill=WHITE, line=None, lw=0.75, radius=0.16, shape=MSO_SHAPE.ROUNDED_RECTANGLE, alpha=None):
    sh = s.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            sf = sh.fill._xPr.find(qn("a:solidFill"))
            clr = sf[0]
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 100000)))
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sh


def line(s, x1, y1, x2, y2, color=PURP, w=1.0, arrow=False, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, In(x1), In(y1), In(x2), In(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", dash)
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


_TOKEN = re.compile(r"(\*\*[^*]+\*\*|==[^=]+==)")


def _add_runs(p, text, size, color, bold, em, em2, italic=False):
    lines = text.split("\n")
    for li, ln_text in enumerate(lines):
        if li:
            p.add_line_break()
        _add_runs_line(p, ln_text, size, color, bold, em, em2, italic)


def _add_runs_line(p, text, size, color, bold, em, em2, italic=False):
    for part in [t for t in _TOKEN.split(text) if t]:
        r = p.add_run()
        f = r.font
        f.size = Pt(size)
        if part.startswith("**"):
            r.text = part[2:-2]
            f.bold = True
            f.color.rgb = rgb(em)
        elif part.startswith("=="):
            r.text = part[2:-2]
            f.bold = True
            f.color.rgb = rgb(em2)
        else:
            r.text = part
            f.bold = bold
            f.color.rgb = rgb(color)
        if italic:
            f.italic = True


def _bullet(p, color=PINK, indent=0.17, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(In(indent))))
    pPr.set("indent", str(-int(In(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    sc = etree.SubElement(buClr, qn("a:srgbClr"))
    sc.set("val", color)
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", char)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    etree.SubElement(pPr, qn("a:buNone"))


def fill_tf(tf, paras, size=12, color=INK, bold=False, em=PURP, em2=PINK, align="l",
            bullets=False, bullet_color=PINK, space_after=0, space_before=0, line_spacing=None, italic=False):
    """paras: список строк или (строка, dict опций). В строке **...** — акцент em, ==...== — акцент em2."""
    tf.word_wrap = True
    first = True
    for item in paras:
        if isinstance(item, tuple):
            txt, o = item
        else:
            txt, o = item, {}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        al = o.get("align", align)
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
        sa = o.get("space_after", space_after)
        sb = o.get("space_before", space_before)
        if sa:
            p.space_after = Pt(sa)
        if sb:
            p.space_before = Pt(sb)
        ls = o.get("line_spacing", line_spacing)
        if ls:
            p.line_spacing = ls
        if o.get("bullet", bullets):
            _bullet(p, o.get("bullet_color", bullet_color))
        else:
            _no_bullet(p)
        _add_runs(p, txt, o.get("size", size), o.get("color", color), o.get("bold", bold),
                  o.get("em", em), o.get("em2", em2), o.get("italic", italic))
    return tf


def tb(s, x, y, w, h, paras, anchor="t", margin=0.0, **kw):
    t = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = t.text_frame
    tf.margin_left = tf.margin_right = In(margin)
    tf.margin_top = tf.margin_bottom = In(0)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for e in list(bodyPr):
        bodyPr.remove(e)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    fill_tf(tf, paras if isinstance(paras, list) else [paras], **kw)
    return t


def shape_text(sh, paras, anchor="m", margin=0.08, **kw):
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = In(margin)
    tf.margin_top = tf.margin_bottom = In(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras if isinstance(paras, list) else [paras], **kw)
    return sh


def box(s, x, y, w, h, paras, fill=WHITE, line=None, radius=0.16, anchor="m", margin=0.1, lw=0.75, **kw):
    sh = rrect(s, x, y, w, h, fill=fill, line=line, radius=radius, lw=lw)
    return shape_text(sh, paras, anchor=anchor, margin=margin, **kw)


def image_dims(path):
    with Image.open(path) as im:
        return im.size


def pic_contain(s, path, x, y, w, h, align="c"):
    pw, ph = image_dims(path)
    k = min(w / pw, h / ph)
    iw, ih = pw * k, ph * k
    if align == "c":
        ix = x + (w - iw) / 2
    elif align == "l":
        ix = x
    else:
        ix = x + w - iw
    return s.shapes.add_picture(path, In(ix), In(y + (h - ih) / 2), In(iw), In(ih))


def pic_cover(s, path, x, y, w, h, fy=0.5, fx=0.5):
    """Картинка заполняет рамку целиком, лишнее обрезается (fx, fy — куда смещать окно)."""
    pw, ph = image_dims(path)
    p = s.shapes.add_picture(path, In(x), In(y), In(w), In(h))
    ar_box, ar_img = w / h, pw / ph
    if ar_img > ar_box:
        keep = ar_box / ar_img
        p.crop_left = (1 - keep) * fx
        p.crop_right = (1 - keep) * (1 - fx)
    else:
        keep = ar_img / ar_box
        p.crop_top = (1 - keep) * fy
        p.crop_bottom = (1 - keep) * (1 - fy)
    return p


def framed(s, path, x, y, w, h, frame="0B0B10", pad=0.05, radius=0.12):
    rrect(s, x, y, w, h, fill=frame, radius=radius)
    return pic_contain(s, path, x + pad, y + pad, w - 2 * pad, h - 2 * pad)


def tag(s, text, x, y, w, color, size=10, h=0.28, tcolor=WHITE):
    return box(s, x, y, w, h, [text], fill=color, radius=h / 2, margin=0.04, size=size, bold=True,
               color=tcolor, align="c")


def set_bg(slide, image_part):
    rId = slide.part.relate_to(image_part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
    cSld = slide._element.cSld
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    xml = (
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="%s"><a:lum/></a:blip><a:srcRect/>'
        '<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>' % rId
    )
    cSld.insert(0, etree.fromstring(xml))


def set_text_keep_format(para, text):
    """Меняет текст абзаца, сохраняя формат первого прогона."""
    runs = para.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def blip_fill_shape(slide, sh, img_path):
    """Заливает фигуру-заполнитель картинкой (картинку заранее обрезать под пропорции фигуры)."""
    img_part, rId = slide.part.get_or_add_image_part(img_path)
    spPr = sh._element.spPr
    for tagname in ("a:noFill", "a:solidFill", "a:blipFill", "a:gradFill"):
        for e in spPr.findall(qn(tagname)):
            spPr.remove(e)
    bf = etree.fromstring(
        '<a:blipFill xmlns:a="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'rotWithShape="1"><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>' % (A_NS, rId)
    )
    # blipFill должен идти после xfrm/prstGeom
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    if geom is None:
        geom = etree.SubElement(spPr, qn("a:prstGeom"))
        geom.set("prst", "rect")
        etree.SubElement(geom, qn("a:avLst"))
    geom.addnext(bf)
