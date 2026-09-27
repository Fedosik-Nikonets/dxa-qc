#!/usr/bin/env python3
"""Сверка прогона перед заморозкой кода с эталоном (27 сентября).

Проверяет каталог результатов `run.sh` (`result.xlsx`, `diagnostics.csv`,
`overlays.zip`) — то, что получит организатор, запустив образ:

1. таблица `result.xlsx`: первые восемь колонок — ровно ТЗ 2.5 в нужном
   порядке, строка на каждый файл, статус только Success/Failure, у Failure —
   причина и пустой `quality_class`, у Success балл `quality_score` в [0, 1]
   больше 0.5 ровно при `quality_class = 1`;
2. `overlays.zip`: PNG и DICOM на каждый успешно обработанный файл;
3. с эталоном (`--reference`, по умолчанию `reports/batch_diagnostics_train.csv`):
   на каждом файле эталона совпадают область, `quality_class`,
   `violation_type` и статус. Ключ — `path_to_study` без обёртки каталога
   (`Исследования/`), которую оставляет прогон архива организатора;
4. численные признаки — максимальное расхождение с эталоном, для сведения:
   другая архитектура и другие версии библиотек не должны менять решения,
   признаки могут расходиться в последних знаках.

`--no-reference` — только пункты 1–2 (стресс-корпус, «Для теста.zip»);
`--no-overlays` — без пункта 2 (прогон без overlay-серии);
`--expect-rows N` — сколько строк должно быть в таблице.

Код возврата 0 — всё сошлось, 1 — есть расхождения (печатаются).

    python scripts/25_freeze_compare.py --out /tmp/freeze/out_train --expect-rows 500
    python scripts/25_freeze_compare.py --out /tmp/freeze/out_stress --no-reference
"""
from __future__ import annotations

import argparse
import csv
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from dxa_qc.report import COLUMNS, STATUS_FAILURE, STATUS_SUCCESS  # noqa: E402

DECISION = ("anatomical_region", "quality_class", "violation_type", "processing_status")
NUMERIC = ("spine_angle_deg", "spine_coverage", "artifact_prob", "femur_field_h_mm", "rotation_value")


def read_xlsx(path: Path) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook  # noqa: PLC0415 — нужен только здесь
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    return header, [{h: ("" if v is None else str(v)) for h, v in zip(header, r)} for r in rows]


def read_csv(path: Path) -> list[dict]:
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def key(path: str, known: set[str]) -> str:
    """path_to_study без обёрток: отрезаем ведущие каталоги, пока не найдём в эталоне."""
    p = path.replace("\\", "/")
    while p not in known and "/" in p:
        p = p.split("/", 1)[1]
    return p


def check_table(out: Path, expect_rows: int | None, problems: list[str]) -> list[dict]:
    header, rows = read_xlsx(out / "result.xlsx")
    if header[:8] != COLUMNS:
        problems.append(f"result.xlsx: первые колонки {header[:8]} вместо {COLUMNS}")
    if expect_rows is not None and len(rows) != expect_rows:
        problems.append(f"result.xlsx: {len(rows)} строк, ожидалось {expect_rows}")
    for r in rows:
        st = r["processing_status"]
        if st not in (STATUS_SUCCESS, STATUS_FAILURE):
            problems.append(f"статус «{st}»: {r['path_to_study']}")
        elif st == STATUS_FAILURE and (r["quality_class"] != "" or not r["violation_type"]):
            problems.append(f"Failure без причины или с классом: {r['path_to_study']}")
        elif st == STATUS_SUCCESS and "quality_score" in r:
            # Балл обязан быть в [0, 1] и больше 0.5 ровно тогда, когда есть нарушение.
            try:
                sc, qc = float(r["quality_score"]), int(float(r["quality_class"]))
            except (TypeError, ValueError):
                problems.append(f"quality_score не число: {r['path_to_study']}")
                continue
            if not 0.0 <= sc <= 1.0 or (sc > 0.5) != (qc == 1):
                problems.append(f"quality_score {sc} противоречит quality_class {qc}: {r['path_to_study']}")
    return rows


def check_overlays(out: Path, rows: list[dict], problems: list[str]) -> int:
    z = out / "overlays.zip"
    if not z.exists():
        problems.append("нет overlays.zip")
        return 0
    names = set(zipfile.ZipFile(z).namelist())
    ok = [r["path_to_study"] for r in rows if r["processing_status"] == STATUS_SUCCESS]
    missing = [p for p in ok if f"{p}.png" not in names or f"{p}.overlay.dcm" not in names]
    for p in missing[:10]:
        problems.append(f"overlay нет: {p}")
    if len(missing) > 10:
        problems.append(f"... и ещё {len(missing) - 10} без overlay")
    return len(ok) - len(missing)


def check_reference(out: Path, ref_path: Path, problems: list[str]) -> dict:
    ref = {r["path_to_study"]: r for r in read_csv(ref_path)}
    got = {}
    for r in read_csv(out / "diagnostics.csv"):
        got[key(r["path_to_study"], set(ref))] = r
    absent = sorted(set(ref) - set(got))
    for p in absent[:10]:
        problems.append(f"нет в прогоне: {p}")
    diff = 0
    for p, e in ref.items():
        g = got.get(p)
        if g is None:
            continue
        bad = [c for c in DECISION if (g.get(c) or "") != (e.get(c) or "")]
        if bad:
            diff += 1
            if diff <= 10:
                problems.append(f"решение разошлось ({', '.join(bad)}): {p}: "
                                + "; ".join(f"{c} {e.get(c)!r} → {g.get(c)!r}" for c in bad))
    if diff > 10:
        problems.append(f"... всего разошлось решений: {diff}")
    dev = {}
    for c in NUMERIC:
        m = 0.0
        for p, e in ref.items():
            g = got.get(p)
            if g is None or not e.get(c) or not g.get(c):
                continue
            m = max(m, abs(float(g[c]) - float(e[c])))
        dev[c] = m
    return {"compared": len(ref) - len(absent), "decision_mismatches": diff, "max_abs_dev": dev}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--out", required=True, type=Path, help="каталог результатов run.sh")
    ap.add_argument("--reference", type=Path, default=ROOT / "reports/batch_diagnostics_train.csv")
    ap.add_argument("--no-reference", action="store_true")
    ap.add_argument("--no-overlays", action="store_true")
    ap.add_argument("--expect-rows", type=int)
    args = ap.parse_args(argv)

    problems: list[str] = []
    rows = check_table(args.out, args.expect_rows, problems)
    status = {s: sum(r["processing_status"] == s for r in rows) for s in (STATUS_SUCCESS, STATUS_FAILURE)}
    print(f"таблица: {len(rows)} строк, {status}")
    if not args.no_overlays:
        print(f"overlay у {check_overlays(args.out, rows, problems)} файлов")
    if not args.no_reference:
        rep = check_reference(args.out, args.reference, problems)
        print(f"эталон {args.reference.name}: сверено {rep['compared']} файлов, "
              f"решений разошлось {rep['decision_mismatches']}")
        print("макс. расхождение признаков: "
              + ", ".join(f"{c} {v:.2g}" for c, v in rep["max_abs_dev"].items()))
    for p in problems:
        print("  ✗", p)
    print("ИТОГ:", "OK" if not problems else f"расхождений {len(problems)}")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
